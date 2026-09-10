"""员工管理路由

提供员工 CRUD 和企微通讯录同步功能。
"""

import io
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession
from uuid import uuid4

from app.database import get_db
from app.models.employee import Employee, EmployeeStatus
from app.routers.admin_auth import get_current_admin, get_current_admin_or_boss
from app.schemas import (
    EmployeeCreateRequest,
    EmployeeUpdateRequest,
    EmployeeResponse,
    EmployeeSyncResult,
)
from app.services.wecom_contact_service import get_contact_service
from app.services.auth_service import hash_password

router = APIRouter(dependencies=[Depends(get_current_admin_or_boss)])


# 批量导入模板的字段定义（顺序即 Excel 列顺序）
# (字段名, 是否必填, 说明)
BATCH_FIELDS: list[tuple[str, bool, str]] = [
    ("姓名", True, "员工真实姓名，必填"),
    ("工号", False, "如 EMP-001，留空时自动生成"),
    ("部门", False, "如 技术部"),
    ("职务", False, "如 工程师"),
    ("手机号", False, "如 13800001234"),
    ("邮箱", False, "如 name@company.com"),
    ("企微ID", False, "企微 UserID，留空时自动生成 manual_xxx"),
]


@router.get("", response_model=list[EmployeeResponse])
async def list_employees(
    keyword: str | None = Query(None, description="搜索姓名/工号/企微ID"),
    status: int | None = Query(None, description="状态: 1=在职 2=离职"),
    db: AsyncSession = Depends(get_db),
):
    """获取员工列表"""
    query = select(Employee).order_by(Employee.created_at.desc())
    if keyword:
        pattern = f"%{keyword}%"
        query = query.where(
            or_(
                Employee.name.ilike(pattern),
                Employee.employee_no.ilike(pattern),
                Employee.wecom_user_id.ilike(pattern),
                Employee.department.ilike(pattern),
            )
        )
    if status is not None:
        query = query.where(Employee.status == EmployeeStatus(status))
    result = await db.execute(query)
    return list(result.scalars().all())


@router.post("/sync", response_model=EmployeeSyncResult)
async def sync_from_wecom(
    db: AsyncSession = Depends(get_db),
    _admin: dict = Depends(get_current_admin),
):
    """从企业微信通讯录同步员工信息

    前置条件：.env 中已配置 WECOM_CORP_ID 和 WECOM_SECRET，
    且企微管理后台为自建应用开通了通讯录读取权限。

    注：企业微信 MCP 功能暂时关闭（main.py 已注释 wecom router 注册）。
    如需恢复，需在 main.py 取消 wecom import 和 router 注册注释，
    并在 docker-compose.yml 启用 wecom-gateway profile。
    本端点当前直接返回 503，不再调企微 API，避免暴露企业密钥或 IP 限制问题。
    """
    from fastapi import HTTPException
    raise HTTPException(
        status_code=503,
        detail="企业微信功能已暂时关闭。如需启用，请取消 main.py 中 wecom router 注释，并在 docker-compose 启用 wecom-gateway profile。",
    )
    service = get_contact_service()
    result = await service.sync(db)
    return result


@router.get("/departments/list", response_model=list[str])
async def list_departments(db: AsyncSession = Depends(get_db)):
    """获取所有部门列表（去重，用于筛选下拉）"""
    from sqlalchemy import distinct, func
    result = await db.execute(
        select(Employee.department)
        .where(Employee.department.isnot(None), Employee.department != "")
        .group_by(Employee.department)
        .order_by(Employee.department)
    )
    return [r[0] for r in result.all()]


@router.get("/{employee_id}", response_model=EmployeeResponse)
async def get_employee(employee_id: int, db: AsyncSession = Depends(get_db)):
    """获取员工详情"""
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalars().first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    return emp


@router.post("", response_model=EmployeeResponse)
async def create_employee(
    request: EmployeeCreateRequest,
    db: AsyncSession = Depends(get_db),
    _admin: dict = Depends(get_current_admin),
):
    """手动添加员工"""
    return await _create_employee_record(request, db)


async def _create_employee_record(
    request: EmployeeCreateRequest, db: AsyncSession
) -> Employee:
    """创建员工记录的内部复用函数

    被 create_employee 和 batch_upload 共用。
    wecom_user_id 为空时自动生成 manual_xxx 前缀。
    """
    # 手动创建时自动生成 wecom_user_id
    wecom_user_id = request.wecom_user_id
    if not wecom_user_id:
        prefix = f"manual_{request.employee_no}_" if request.employee_no else "manual_"
        wecom_user_id = f"{prefix}{uuid4().hex[:8]}"

    # 检查 wecom_user_id 唯一性
    existing = await db.execute(
        select(Employee).where(Employee.wecom_user_id == wecom_user_id)
    )
    if existing.scalars().first():
        raise HTTPException(
            status_code=409,
            detail=f"wecom_user_id '{wecom_user_id}' 已存在",
        )

    emp = Employee(
        wecom_user_id=wecom_user_id,
        name=request.name,
        employee_no=request.employee_no,
        department=request.department,
        department_id=request.department_id,
        position=request.position,
        mobile=request.mobile,
        email=request.email,
        status=EmployeeStatus(request.status),
    )

    # 如果请求中有密码字段，设置密码
    if request.password:
        emp.password_hash = hash_password(request.password)

    db.add(emp)
    await db.commit()
    await db.refresh(emp)
    return emp


@router.get("/batch/template")
async def download_batch_template():
    """下载批量添加员工的 Excel 模板

    模板含 2 个 sheet：
    1. 填写表：表头 + 2 行示例数据（可直接修改后上传）
    2. 字段说明：每个字段的含义、是否必填、示例
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # Sheet 1: 填写表
    ws = wb.active
    ws.title = "员工信息"
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_idx, (name, _, _) in enumerate(BATCH_FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 示例数据（2 行）
    samples = [
        ["张三", "EMP-001", "技术部", "工程师", "13800001234", "zhangsan@company.com", ""],
        ["李四", "EMP-002", "销售部", "经理", "13800005678", "lisi@company.com", "wecom_lisi"],
    ]
    for row_idx, sample in enumerate(samples, 2):
        for col_idx, value in enumerate(sample, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # 列宽自适应
    for col_idx, (name, _, _) in enumerate(BATCH_FIELDS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(name) * 2 + 4)

    # Sheet 2: 字段说明
    ws2 = wb.create_sheet("字段说明")
    ws2.append(["字段名", "是否必填", "说明"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for name, required, desc in BATCH_FIELDS:
        ws2.append([name, "必填" if required else "选填", desc])
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 50

    # 导出为字节流
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employee_batch_template.xlsx"},
    )


@router.post("/batch/upload")
async def batch_upload_employees(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _admin: dict = Depends(get_current_admin),
):
    """批量上传 Excel 添加员工

    接收 .xlsx 文件，按模板字段顺序解析每行，循环创建员工记录。
    返回 {total, success, failed, errors} 统计。
    """
    from openpyxl import load_workbook

    # 读取上传文件
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件为空")

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Excel 文件解析失败：{e}。请使用模板填写后上传。",
        )

    # 取第一个 sheet
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="文件无数据行（仅表头）")

    # 第一行是表头，跳过；从第二行开始处理
    data_rows = rows[1:]
    total = 0
    success = 0
    failed = 0
    errors: list[str] = []

    for i, row in enumerate(data_rows, 1):
        # 跳过完全空白的行
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        total += 1

        # 按 BATCH_FIELDS 顺序取值
        values = list(row) + [None] * len(BATCH_FIELDS)  # 补齐列数
        name = str(values[0]).strip() if values[0] is not None else ""
        employee_no = str(values[1]).strip() if values[1] is not None else ""
        department = str(values[2]).strip() if values[2] is not None else ""
        position = str(values[3]).strip() if values[3] is not None else ""
        mobile = str(values[4]).strip() if values[4] is not None else ""
        email = str(values[5]).strip() if values[5] is not None else ""
        wecom_user_id = str(values[6]).strip() if values[6] is not None else ""

        if not name:
            failed += 1
            errors.append(f"第 {i} 行：姓名为空")
            continue

        try:
            await _create_employee_record(
                EmployeeCreateRequest(
                    name=name,
                    employee_no=employee_no or None,
                    department=department or None,
                    position=position or None,
                    mobile=mobile or None,
                    email=email or None,
                    wecom_user_id=wecom_user_id or None,
                    status=1,
                ),
                db,
            )
            success += 1
        except HTTPException as e:
            failed += 1
            errors.append(f"第 {i} 行 {name}：{e.detail}")
        except Exception as e:
            failed += 1
            errors.append(f"第 {i} 行 {name}：{str(e)[:100]}")

    return {
        "total": total,
        "success": success,
        "failed": failed,
        "errors": errors[:50],  # 最多返回前 50 条错误
    }


@router.put("/{employee_id}", response_model=EmployeeResponse)
async def update_employee(
    employee_id: int,
    request: EmployeeUpdateRequest,
    db: AsyncSession = Depends(get_db),
    _admin: dict = Depends(get_current_admin),
):
    """更新员工信息"""
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalars().first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    update_data = request.model_dump(exclude_unset=True)
    # 处理密码更新
    password = update_data.pop('password', None)
    if password:
        emp.password_hash = hash_password(password)
    for field, value in update_data.items():
        if field == "status":
            emp.status = EmployeeStatus(value)
        else:
            setattr(emp, field, value)

    await db.commit()
    await db.refresh(emp)
    return emp


@router.delete("/{employee_id}")
async def delete_employee(
    employee_id: int,
    db: AsyncSession = Depends(get_db),
    _admin: dict = Depends(get_current_admin),
):
    """删除员工"""
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalars().first()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")

    await db.delete(emp)
    await db.commit()
    return {"message": f"员工 {emp.name} 已删除", "deleted": True}
